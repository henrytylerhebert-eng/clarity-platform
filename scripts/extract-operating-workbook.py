#!/usr/bin/env python3
"""Read the accepted synthetic workbook into versioned runtime records; never writes Excel."""
import argparse
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries

TABLES = [
    ("patients", "PATIENTS", "tPatients", False),
    ("payers", "PAYERS", "tPayers", False),
    ("services", "SERVICES", "tServices", False),
    ("contractRates", "CONTRACT_RATES", "tRates", False),
    ("encounters", "ENCOUNTERS", "tEnc", False),
    ("inpatient", "INPATIENT", "tIP", False),
    ("iopRoster", "IOP_ROSTER", "tEnroll", False),
    ("iopVisits", "IOP_VISITS", "tVisits", False),
    ("iopSessions", "IOP_SESSIONS", "tSessions", False),
    ("sessionAttendance", "SESSION_ATTENDANCE", "tSessionAtt", False),
    ("serviceLedger", "SERVICE_LEDGER", "tService", False),
    ("staffStandards", "STAFF_STANDARDS", "tLaborRates", False),
    ("staffDetail", "STAFF_DETAIL", "tLabor", False),
    ("budget", "BUDGET", "tBudget", False),
    ("forecast", "FORECAST", "tForecast", False),
    ("forecastMix", "FORECAST", "tForecastMix", False),
    ("ancillary", "ANCILLARY", "tAnc", False),
    ("collections", "COLLECTIONS", "tCash", False),
    ("receiptAllocations", "RECEIPT_ALLOCATIONS", "tAlloc", False),
    ("invoices", "INVOICES", "tInvoices", False),
    ("urPayer", "UR_PAYER", "tUR", False),
    ("coverage", "COVERAGE", "tCoverage", False),
    ("monthlySnapshot", "MONTHLY", "tMonthly", True),
    ("iopSnapshot", "IOP", "tIOP", True),
    ("staffingSnapshot", "STAFFING", "tStaff", True),
    ("programBudgetSnapshot", "BUDGET", "tProgramBudget", True),
    ("managementSnapshot", "MANAGEMENT_REPORT", "tManagement", True),
    ("assistanceSnapshot", "ASSISTANCE", "tAssistance", True),
    ("benefitSnapshot", "BENEFIT_REVIEW", "tBenefitReview", True),
    ("payerMixSnapshot", "PAYER_MIX", "tMix", True),
    ("stateMixSnapshot", "STATE_MIX", "tStateMix", True),
]

EXPLICIT_NUMBERS = {
    "staffDetail": {"productiveHours", "agencyHours", "oneToOneHours", "ptoHours", "trainingHours"},
    "iopVisits": {"servicesDelivered", "meals", "groupServices", "nonpayableMeals"},
    "iopSessions": {"countSession"}, "sessionAttendance": {"units"},
    "serviceLedger": {"units", "grossCharges", "signedAllowanceAdjustment"},
    "invoices": {"signedAdjustment", "vendorPaid"},
    "urPayer": {"submittedDays", "authorizedDays", "deniedDays", "pendingDays"},
}

REGISTRY_REQUIRED = {
    "payers": {"payerId", "payerAdministrator", "planProduct", "lineOfBusiness", "network", "funding", "active"},
    "services": {"serviceId", "program", "setting", "code", "billableUnit", "description", "pricingDateBasis"},
    "contractRates": {"rateId", "payerId", "serviceId", "method", "rateValue", "effectiveFrom", "effectiveTo", "status", "evidence", "reviewer", "version", "pricingBasis", "scope"},
}


def key(label):
    words = re.findall(r"[A-Za-z0-9]+", label)
    return words[0].lower() + "".join(w.title() for w in words[1:])


def cell_value(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if value == "":
        return None
    if not isinstance(value, (str, int, float, bool, type(None))):
        raise ValueError(f"Unsupported cached value {type(value)}")
    return value


def extract(source):
    formulas = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True, read_only=True)
    output = {
        "schemaVersion": 1,
        "source": {
            "name": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "year": 2026,
            "importedAt": dt.datetime.now(dt.timezone.utc).isoformat(),
        },
        "settings": {
            "hospital": cached["SETUP"]["B5"].value,
            "licensedBeds": cached["SETUP"]["B9"].value,
            "reportingCutoff": cell_value(cached["SETUP"]["B12"].value),
        },
        "tables": [],
    }
    for table_key, sheet_name, table_name, snapshot in TABLES:
        sheet = formulas[sheet_name]
        left, top, right, bottom = range_boundaries(sheet.tables[table_name].ref)
        raw_rows = list(sheet.iter_rows(min_row=top + 1, max_row=bottom, min_col=left, max_col=right))
        cached_rows = list(cached[sheet_name].iter_rows(min_row=top + 1, max_row=bottom, min_col=left, max_col=right, values_only=True))
        labels = [str(sheet.cell(top, c).value) for c in range(left, right + 1)]
        keys = [key(label) for label in labels]
        active = [(i, raw, values) for i, (raw, values) in enumerate(zip(raw_rows, cached_rows)) if values[0] not in (None, "")]
        columns = []
        for j, (label, column_key) in enumerate(zip(labels, keys)):
            cells = [raw[j] for _, raw, _ in active]
            has_formula = any(c.data_type == "f" for c in cells)
            has_input = any(c.data_type != "f" and c.value is not None for c in cells)
            date_type = any(isinstance(c.value, (dt.datetime, dt.date)) for c in cells)
            if not date_type and (label.lower().endswith("date") or label in ("Effective from", "Effective to", "Period start", "Period end", "Review due")):
                date_type = True
            number_type = any(isinstance(values[j], (int, float)) and not isinstance(values[j], bool) for _, _, values in active)
            if label in ("Opening override", "Patient share input"):
                number_type = True  # Intentionally blank numeric inputs must remain usable.
            editable = not snapshot and (not has_formula or has_input)
            if j == 0 or table_key in ("collections", "receiptAllocations"):
                editable = False  # identities and signed receipt history are immutable
            col = {"key": column_key, "label": label, "type": "date" if date_type else "number" if number_type else "text", "editable": editable}
            if table_key == "contractRates" and column_key == "rateValue":
                col["label"] = "Rate value (USD/unit or fraction 0–1)"
            if snapshot or has_formula:
                col["snapshot"] = True
            if column_key.startswith("signed") or (table_key == "receiptAllocations" and column_key == "amount"):
                col["signed"] = True
            if column_key in ("patientId", "payerId", "serviceId", "evidence", "admitDate", "serviceDate", "date", "effectiveFrom", "effectiveTo", "role", "program", "attendance", "status"):
                col["required"] = True
            if column_key in EXPLICIT_NUMBERS.get(table_key, set()):
                col["required"] = True
            if column_key in REGISTRY_REQUIRED.get(table_key, set()):
                col["required"] = True
            options = {"attendance": ["ATTENDED", "NO SHOW", "CANCELLED"], "rateMethod": ["BLENDED", "PAYER MIX"], "method": ["PER UNIT", "PCT CHARGES"]}.get(column_key)
            if options:
                col["options"] = options
            columns.append(col)
        rows = []
        for i, raw, values in active:
            row = {"id": f"{table_key}:{top + i + 1}", "sourceRow": top + i + 1, "values": {k: cell_value(v) for k, v in zip(keys, values)}}
            formula_keys = [keys[j] for j, c in enumerate(raw) if c.data_type == "f"]
            if formula_keys:
                row["formulaKeys"] = formula_keys
            rows.append(row)
        table = {"key": table_key, "title": str(sheet["A1"].value), "sheet": sheet_name, "columns": columns, "rows": rows}
        date_key = next((k for k in ("date", "serviceDate", "receiptDate", "admitDate", "startDate", "periodStart") if k in keys), None)
        if date_key:
            table["dateKey"] = date_key
        if snapshot:
            table["snapshot"] = True
        output["tables"].append(table)
    formulas.close()
    cached.close()
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    data = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(json.dumps({"sha256": data["source"]["sha256"], "tables": {t["key"]: len(t["rows"]) for t in data["tables"]}}, indent=2))
